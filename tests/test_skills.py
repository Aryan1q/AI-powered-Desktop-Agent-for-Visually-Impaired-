"""A round trip must be right, and must fail loudly rather than quietly.

The dangerous failure is not a crash: it is a PDF that comes back looking fine
while the edit silently did nothing, or while the text landed somewhere the
user never asked for. These tests pin that down without touching the network.
"""
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from tests._support import ROOT  # noqa: F401


def make_docx(path: Path, paragraphs, table=None):
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        rows = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                rows.cell(r, c).text = value
    document.save(str(path))
    return path


def read_docx(path: Path):
    import docx

    document = docx.Document(str(path))
    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs)
    return texts


class TestEditDocx(unittest.TestCase):
    def test_replaces_in_paragraphs(self):
        from ufo.skills.pdf_roundtrip import edit_docx

        with TemporaryDirectory() as d:
            path = make_docx(Path(d) / "a.docx",
                             ["Issue year: 2024", "no match here"])
            count = edit_docx(path, "2024", "2025")
            self.assertEqual(count, 1)
            self.assertIn("Issue year: 2025", read_docx(path))

    def test_counts_every_occurrence(self):
        from ufo.skills.pdf_roundtrip import edit_docx

        with TemporaryDirectory() as d:
            path = make_docx(Path(d) / "a.docx",
                             ["2024", "2024", "nothing"])
            self.assertEqual(edit_docx(path, "2024", "2025"), 2)

    def test_reaches_into_tables(self):
        from ufo.skills.pdf_roundtrip import edit_docx

        with TemporaryDirectory() as d:
            path = make_docx(Path(d) / "a.docx", ["header"],
                             table=[["year", "2024"], ["other", "x"]])
            self.assertEqual(edit_docx(path, "2024", "2025"), 1)
            self.assertIn("2025", read_docx(path))

    def test_text_split_across_runs_is_still_found(self):
        """A converter routinely splits one sentence into several runs.

        Matching run by run would miss it, and the user's edit would appear to
        be ignored.
        """
        import docx

        from ufo.skills.pdf_roundtrip import edit_docx

        with TemporaryDirectory() as d:
            path = Path(d) / "a.docx"
            document = docx.Document()
            paragraph = document.add_paragraph()
            for piece in ("Issue year: ", "20", "24", " end"):
                paragraph.add_run(piece)
            document.save(str(path))
            self.assertEqual(edit_docx(path, "2024", "2025"), 1)
            self.assertIn("Issue year: 2025 end", read_docx(path))

    def test_regex_metacharacters_are_literal(self):
        """A document full of C++ and A.B.C must not be treated as a pattern."""
        from ufo.skills.pdf_roundtrip import edit_docx

        with TemporaryDirectory() as d:
            path = make_docx(Path(d) / "a.docx",
                             ["Version A.B.C released", "no dots here"])
            self.assertEqual(edit_docx(path, "A.B.C", "A.X.C"), 1)
            self.assertIn("Version A.X.C released", read_docx(path))

    def test_nothing_found_is_an_error_not_a_silent_success(self):
        """The most important case: report it, do not pretend it worked."""
        from ufo.skills.pdf_roundtrip import RoundtripError, edit_docx

        with TemporaryDirectory() as d:
            path = make_docx(Path(d) / "a.docx", ["nothing relevant"])
            with self.assertRaises(RoundtripError) as ctx:
                edit_docx(path, "2024", "2025")
        self.assertIn("2024", str(ctx.exception))

    def test_case_sensitivity_can_be_required(self):
        from ufo.skills.pdf_roundtrip import RoundtripError, edit_docx

        with TemporaryDirectory() as d:
            path = make_docx(Path(d) / "a.docx", ["total: abc"])
            self.assertEqual(edit_docx(path, "ABC", "XYZ"), 1)
            path2 = make_docx(Path(d) / "b.docx", ["total: abc"])
            with self.assertRaises(RoundtripError):
                edit_docx(path2, "ABC", "XYZ", match_case=True)

    def test_empty_old_text_is_refused(self):
        from ufo.skills.pdf_roundtrip import RoundtripError, edit_docx

        with TemporaryDirectory() as d:
            path = make_docx(Path(d) / "a.docx", ["text"])
            with self.assertRaises(RoundtripError):
                edit_docx(path, "", "x")

    def test_the_file_stays_a_valid_docx(self):
        import docx

        from ufo.skills.pdf_roundtrip import edit_docx

        with TemporaryDirectory() as d:
            path = make_docx(Path(d) / "a.docx", ["year 2024"])
            edit_docx(path, "2024", "2025")
            docx.Document(str(path))          # must not raise


def make_text_pdf(path: Path, text: str = "Issue year: 2024") -> Path:
    """A PDF with a real text layer, so extract_text has something to find."""
    import shutil
    import subprocess

    chrome = next((p for p in (
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        shutil.which("chrome"), shutil.which("google-chrome"),
    ) if p and Path(p).exists()), None)
    if chrome is None:
        raise unittest.SkipTest("no Chrome available to build a text PDF")

    html = path.with_suffix(".html")
    html.write_text(
        f"<!doctype html><meta charset='utf-8'>"
        f"<body style='font-family:Georgia,serif;font-size:14pt'>"
        f"<p>{text}</p></body>", encoding="utf-8")
    try:
        subprocess.run(
            [chrome, "--headless", "--disable-gpu", "--no-pdf-header-footer",
             f"--print-to-pdf={path}", html.as_uri()],
            check=True, capture_output=True, timeout=120)
    except (subprocess.SubprocessError, OSError) as e:
        raise unittest.SkipTest(f"Chrome could not print a PDF: {e}") from e
    if not path.exists():
        raise unittest.SkipTest("Chrome produced no PDF")
    return path


class TestVerifyPdf(unittest.TestCase):
    def test_a_real_pdf_passes(self):
        from ufo.skills.pdf_roundtrip import verify_pdf

        with TemporaryDirectory() as d:
            path = make_text_pdf(Path(d) / "out.pdf", "hello world")
            verify_pdf(path)                     # must not raise

    def test_the_replacement_is_found_in_the_finished_pdf(self):
        from ufo.skills.pdf_roundtrip import verify_pdf

        with TemporaryDirectory() as d:
            path = make_text_pdf(Path(d) / "out.pdf", "Issue year: 2025")
            self.assertIn("2025", verify_pdf(path, "2025"))

    def test_an_edit_that_did_not_survive_is_caught(self):
        """The whole point: a PDF that came back without the edit is a failure."""
        from ufo.skills.pdf_roundtrip import RoundtripError, verify_pdf

        with TemporaryDirectory() as d:
            path = make_text_pdf(Path(d) / "out.pdf", "Issue year: 2024")
            with self.assertRaises(RoundtripError) as ctx:
                verify_pdf(path, "2025")
            self.assertIn("2025", str(ctx.exception))

    def test_missing_file_is_an_error(self):
        from ufo.skills.pdf_roundtrip import RoundtripError, verify_pdf

        with TemporaryDirectory() as d:
            with self.assertRaises(RoundtripError):
                verify_pdf(Path(d) / "nope.pdf")

    def test_a_non_pdf_named_pdf_is_rejected(self):
        """A converter error page saved as .pdf must not pass as a document."""
        from ufo.skills.pdf_roundtrip import RoundtripError, verify_pdf

        with TemporaryDirectory() as d:
            path = Path(d) / "bad.pdf"
            path.write_text("<html>conversion failed</html>", encoding="utf-8")
            with self.assertRaises(RoundtripError) as ctx:
                verify_pdf(path)
            self.assertIn("readable PDF", str(ctx.exception))


class TestSpokenInstruction(unittest.TestCase):
    def test_quoted_text_wins(self):
        from ufo.skills.pdf_roundtrip import find_replacement

        self.assertEqual(
            find_replacement('change "Issue year: 2024" to "Issue year: 2025"'),
            ("Issue year: 2024", "Issue year: 2025"))

    def test_replace_with_wording(self):
        from ufo.skills.pdf_roundtrip import find_replacement

        self.assertEqual(
            find_replacement("replace 2024 with 2025 in the document"),
            ("2024", "2025"))

    def test_change_to_wording(self):
        from ufo.skills.pdf_roundtrip import find_replacement

        self.assertEqual(find_replacement("change 2024 to 2025"),
                         ("2024", "2025"))

    def test_filler_words_are_stripped(self):
        from ufo.skills.pdf_roundtrip import find_replacement

        self.assertEqual(
            find_replacement("replace the text 2024 with the text 2025"),
            ("2024", "2025"))

    def test_no_replacement_returns_none(self):
        from ufo.skills.pdf_roundtrip import find_replacement

        self.assertIsNone(find_replacement("open the report and read it"))
        self.assertIsNone(find_replacement(""))

    def test_a_lone_quote_pair_is_not_used(self):
        """One quoted phrase is not a before-and-after pair."""
        from ufo.skills.pdf_roundtrip import find_replacement

        self.assertIsNone(find_replacement('summarise the "quarterly report"'))


class TestFindPdf(unittest.TestCase):
    def test_exact_filename_wins(self):
        from ufo.skills.pdf_roundtrip import find_pdf

        with TemporaryDirectory() as d:
            target = Path(d) / "Annual_Report.pdf"
            target.write_bytes(b"%PDF-1.4")
            (Path(d) / "other.pdf").write_bytes(b"%PDF-1.4")
            found = find_pdf("convert Annual_Report.pdf to word",
                             [Path(d)])
            self.assertEqual(found.name, "Annual_Report.pdf")

    def test_words_in_the_name_are_matched(self):
        from ufo.skills.pdf_roundtrip import find_pdf

        with TemporaryDirectory() as d:
            (Path(d) / "cloud_computing_assignment.pdf").write_bytes(b"%PDF-1.4")
            found = find_pdf("please take the cloud computing assignment pdf",
                             [Path(d)])
            self.assertIsNotNone(found)
            self.assertIn("cloud", found.name)

    def test_a_weak_match_is_refused(self):
        """Guessing a document is worse than asking which one."""
        from ufo.skills.pdf_roundtrip import find_pdf

        with TemporaryDirectory() as d:
            (Path(d) / "cloud_computing_assignment.pdf").write_bytes(b"%PDF-1.4")
            self.assertIsNone(find_pdf("convert the pdf", [Path(d)]))

    def test_missing_directory_is_skipped(self):
        from ufo.skills.pdf_roundtrip import find_pdf

        self.assertIsNone(find_pdf("anything", [Path("no/such/place")]))


class TestWorkflowGuards(unittest.TestCase):
    def test_a_non_pdf_source_is_refused_before_opening_a_browser(self):
        from ufo.skills.pdf_roundtrip import RoundtripError, roundtrip

        with TemporaryDirectory() as d:
            path = Path(d) / "notes.txt"
            path.write_text("hello", encoding="utf-8")
            with self.assertRaises(RoundtripError) as ctx:
                roundtrip(path, "a", "b", out_dir=d)
        self.assertIn("not a PDF", str(ctx.exception))

    def test_a_missing_source_is_refused(self):
        from ufo.skills.pdf_roundtrip import RoundtripError, roundtrip

        with TemporaryDirectory() as d:
            with self.assertRaises(RoundtripError):
                roundtrip(Path(d) / "absent.pdf", "a", "b", out_dir=d)

    def test_the_source_pdf_is_never_overwritten(self):
        """The original is the user's only copy; the result gets a new name."""
        import ufo.skills.pdf_roundtrip as skill

        with TemporaryDirectory() as d:
            directory = Path(d)
            source = directory / "report.pdf"
            source.write_bytes(b"%PDF-1.4 original")
            before = source.read_bytes()

            class FakeDriver:
                def get(self, url):
                    pass

                def execute_script(self, *a):
                    pass

            # Fail at the very first real step: the file must be untouched.
            with self.assertRaises(Exception):
                skill.convert_file(FakeDriver(), source,
                                   skill.ILOVEPDF_PDF_TO_WORD, directory / "dl",
                                   "pdf to word")
            self.assertEqual(source.read_bytes(), before)

    def test_repeated_runs_do_not_overwrite_the_previous_result(self):
        import ufo.skills.pdf_roundtrip as skill

        with TemporaryDirectory() as d:
            directory = Path(d)
            (directory / "report_edited.pdf").write_bytes(b"%PDF-1.4 one")
            self.assertEqual(
                skill._unique_output(directory, "report").name,
                "report_edited_2.pdf")
